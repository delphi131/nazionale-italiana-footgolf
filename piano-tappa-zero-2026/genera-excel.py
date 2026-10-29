#!/usr/bin/env python3
"""
Genera il piano operativo e il preventivo della Tappa Zero in Excel a partire da `dati.json`.

    python3 piano-tappa-zero-2026/genera-excel.py

Fogli: Attività, Preventivo, Riepilogo per area, Community. Il preventivo resta vivo:
totali, riepiloghi e grafico sono formule, quindi cambiando quantità, prezzi o lo stato
di una voce si aggiorna tutto senza rigenerare.
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DIR = Path(__file__).resolve().parent
DATI = json.loads((DIR / "dati.json").read_text(encoding="utf-8"))
OUT = DIR / "dist" / "piano-tappa-zero-2026.xlsx"
# La numerazione delle voci segue l'ordine del file, come nel PDF
for _i, _v in enumerate(DATI["voci"], start=1):
    _v["n"] = _i

# Token del sito AIFG (design_handoff_aifg_footgolf)
INK, LIGHT, PANEL, BLUE, BLUE_DARK, GOLD = "0A0C0A", "EEF1EE", "F4F6F4", "0072CE", "0A5FA0", "E8C25A"
GRIGIO, LINEA, INPUT = "6B706B", "D5DAD5", "FFF8E1"

FONT = "Arial"
EUR = '€ #,##0.00;-€ #,##0.00;"–"'
EUR_UNIT = '€ #,##0.00##;-€ #,##0.00##;"–"'

def f(size=10, bold=False, color=INK, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)

def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)

sottile = Side(style="thin", color=LINEA)
BORDO = Border(top=sottile, bottom=sottile, left=sottile, right=sottile)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
DESTRA = Alignment(horizontal="right", vertical="center")

wb = Workbook()
ws = wb.active
ws.title = "Preventivo"
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = BLUE

# ---------------------------------------------------------------- colonne
COLONNE = [
    # (intestazione, larghezza)
    ("N.", 5),
    ("Area", 20),
    ("Voce", 26),
    ("Descrizione / specifiche", 38),
    ("Fornitore", 15),
    ("Tempi", 16),
    ("Q.tà", 8),
    ("Unità", 10),
    ("Prezzo unitario", 14),
    ("Totale", 14),
    ("IVA", 10),
    ("Stato", 13),
    ("Nel totale", 10),
    ("Ordine", 13),
    ("Link", 24),
    ("Foto", 22),
    ("Note", 36),
]
C = {nome: get_column_letter(i + 1) for i, (nome, _) in enumerate(COLONNE)}
ULTIMA = get_column_letter(len(COLONNE))
for i, (_, w) in enumerate(COLONNE, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

# ---------------------------------------------------------------- testata
for r in range(1, 4):
    for c in range(1, len(COLONNE) + 1):
        ws.cell(r, c).fill = fill(INK)
ws.row_dimensions[1].height = 30
ws.row_dimensions[2].height = 22
ws.row_dimensions[3].height = 18

logo = XLImage(str(DIR / "assets" / "aifg-logo.png"))
logo.width, logo.height = 64, 63
ws.add_image(logo, "A1")

ev = DATI["evento"]
ws["C1"] = f"PREVENTIVO DI SPESA · {ev['titolo'].upper()}"
ws["C1"].font = f(18, True, "FFFFFF")
ws["C1"].alignment = Alignment(vertical="bottom")
ws["C2"] = f"{ev['date']}  ·  {ev['luogo']}, {ev['citta']}"
ws["C2"].font = f(11, False, "4FB3F0")
ws["C3"] = f"Associazione Italiana FootGolf  ·  materiale, allestimenti, staff e ospitalità  ·  aggiornato al {ev['aggiornato']}"
ws["C3"].font = f(8.5, False, "9DA39D")

# ---------------------------------------------------------------- tabella voci
R_HEAD = 10
R0 = R_HEAD + 1
VOCI = DATI["voci"]
RIGHE_LIBERE = 5
R_LAST = R0 + len(VOCI) + RIGHE_LIBERE - 1
R_TOT = R_LAST + 1

def rng(col):
    return f"${C[col]}${R0}:${C[col]}${R_LAST}"

# ---------------------------------------------------------------- KPI
KPI = [
    ("B", "C", "Totale base", f'=SUMIFS({rng("Totale")},{rng("Nel totale")},"Sì")', EUR, BLUE),
    # Le alternative (A3 al posto della brochure) non si sommano: si sceglie l'una o l'altra
    ("D", "D", "Opzionali", f'=SUMIFS({rng("Totale")},{rng("Stato")},"Opzionale",{rng("Nel totale")},"No")', EUR, GRIGIO),
    ("E", "G", "Totale con opzionali", "=B6+D6", EUR, INK),
    ("H", "I", "Voci da definire", f'=COUNTIFS({rng("Stato")},"Da definire")', "0", "B7791F"),
    ("J", "L", "Budget disponibile", None, EUR, INK),
    ("M", "N", "Residuo sul totale base", "=IF(J6=\"\",\"\",J6-B6)", EUR, BLUE_DARK),
]
ws.row_dimensions[5].height = 18
ws.row_dimensions[6].height = 34
for c1, c2, etichetta, formula, fmt, colore in KPI:
    ws.merge_cells(f"{c1}5:{c2}5")
    ws.merge_cells(f"{c1}6:{c2}6")
    ws[f"{c1}5"] = etichetta.upper()
    ws[f"{c1}5"].font = f(8, True, GRIGIO)
    ws[f"{c1}5"].alignment = Alignment(vertical="bottom")
    cel = ws[f"{c1}6"]
    cel.value = formula
    cel.number_format = fmt
    cel.font = f(20, True, colore)
    cel.alignment = Alignment(horizontal="left", vertical="center")
    for col in range(ws[f"{c1}6"].column, ws[f"{c2}6"].column + 1):
        ws.cell(6, col).border = Border(bottom=Side(style="medium", color=colore))
        ws.cell(5, col).fill = fill(PANEL)
        ws.cell(6, col).fill = fill(PANEL)
ws["J6"].fill = fill(INPUT)
ws["J6"].comment = Comment("Inserire qui il budget stanziato per la tappa: il residuo si calcola da solo.", "AIFG")

ws["B8"] = "LEGENDA"
ws["B8"].font = f(8, True, GRIGIO)
ws["C8"] = "Celle gialle = da compilare (quantità, prezzi, fornitore, link, foto).  Stato e Nel totale: menu a tendina.  Il Totale si calcola da Q.tà × Prezzo unitario."
ws["C8"].font = f(9, False, GRIGIO, italic=True)

# intestazione
ws.row_dimensions[R_HEAD].height = 26
for i, (nome, _) in enumerate(COLONNE, start=1):
    cel = ws.cell(R_HEAD, i, nome.upper())
    cel.font = f(8.5, True, "FFFFFF")
    cel.fill = fill(BLUE)
    cel.alignment = CENTRO if nome not in ("Voce", "Descrizione / specifiche", "Note") else Alignment(vertical="center")
    cel.border = Border(bottom=Side(style="medium", color=INK))

AREE = {a["id"]: f"{a['id']} · {a['nome']}" for a in DATI["aree"]}

def scrivi(r, col, v, fmt=None, font=None, align=WRAP, input_=False):
    cel = ws[f"{C[col]}{r}"]
    cel.value = v
    cel.font = font or f(10)
    cel.alignment = align
    cel.border = BORDO
    if fmt:
        cel.number_format = fmt
    if input_:
        cel.fill = fill(INPUT)
    return cel

for k in range(len(VOCI) + RIGHE_LIBERE):
    r = R0 + k
    v = VOCI[k] if k < len(VOCI) else None
    ws.row_dimensions[r].height = 78  # spazio per la foto incollata nella cella
    zebra = fill("FFFFFF") if k % 2 == 0 else fill("F8FAF8")
    for col in range(1, len(COLONNE) + 1):
        ws.cell(r, col).fill = zebra

    scrivi(r, "N.", v["n"] if v else k + 1, font=f(11, True, BLUE), align=CENTRO)
    scrivi(r, "Area", AREE[v["area"]] if v else None, font=f(9, False, GRIGIO))
    scrivi(r, "Voce", v["nome"] if v else None, font=f(10.5, True))
    scrivi(r, "Descrizione / specifiche", v["desc"] if v else None, font=f(9.5, False, "3A403A"))
    scrivi(r, "Fornitore", (v.get("fornitore") or None) if v else None, input_=True)
    scrivi(r, "Tempi", (v.get("tempi") or None) if v else None, input_=True, font=f(9, False, "3A403A"))
    scrivi(r, "Q.tà", v["qta"] if v else None, "#,##0", align=CENTRO, input_=True)
    scrivi(r, "Unità", v["unita"] if v else None, align=CENTRO, font=f(9, False, GRIGIO))

    # Prezzo unitario: dove il preventivo dà un importo complessivo lo si ricava,
    # così il Totale torna esatto e il prezzo per pezzo resta leggibile.
    if v and v.get("totale") is not None:
        pu = f"={v['totale']}/{v['qta']}"
        cel = scrivi(r, "Prezzo unitario", pu, EUR_UNIT, align=DESTRA, input_=True)
        cel.comment = Comment(
            f"Dal preventivo: € {v['totale']:.2f} complessivi per {v['qta']} {v['unita']}.".replace(".", ","),
            "AIFG")
    else:
        scrivi(r, "Prezzo unitario", v.get("prezzo") if v else None, EUR_UNIT, align=DESTRA, input_=True)

    q, p = f"{C['Q.tà']}{r}", f"{C['Prezzo unitario']}{r}"
    scrivi(r, "Totale", f'=IF(OR({q}="",{p}=""),"",{q}*{p})', EUR, font=f(11, True), align=DESTRA)
    scrivi(r, "IVA", (v.get("iva") or "Da verificare") if v else None, align=CENTRO, font=f(9, False, GRIGIO))
    scrivi(r, "Stato", v["stato"] if v else None, align=CENTRO, font=f(9.5, True))
    scrivi(r, "Nel totale", ("Sì" if v["nelTotale"] else "No") if v else None, align=CENTRO)
    scrivi(r, "Ordine", "Da ordinare" if v else None, align=CENTRO, font=f(9.5))
    link = scrivi(r, "Link", None, input_=True, font=f(9, False, BLUE, underline="single"))
    if v and v.get("link"):
        link.value = v["link"].replace("https://", "").replace("www.", "")
        link.hyperlink = v["link"]
    foto = scrivi(r, "Foto", "incolla qui la foto", input_=True, align=CENTRO,
                  font=f(8, False, "B9BEB9", italic=True))
    scrivi(r, "Note", (v.get("note") or None) if v else None, font=f(9, False, "3A403A"))

# riga totale
ws.row_dimensions[R_TOT].height = 28
for col in range(1, len(COLONNE) + 1):
    ws.cell(R_TOT, col).fill = fill(INK)
ws[f"{C['Voce']}{R_TOT}"] = "TOTALE BASE (voci «Nel totale» = Sì)"
ws[f"{C['Voce']}{R_TOT}"].font = f(10, True, "FFFFFF")
ws[f"{C['Voce']}{R_TOT}"].alignment = Alignment(vertical="center")
ws[f"{C['Totale']}{R_TOT}"] = f'=SUMIFS({rng("Totale")},{rng("Nel totale")},"Sì")'
ws[f"{C['Totale']}{R_TOT}"].number_format = EUR
ws[f"{C['Totale']}{R_TOT}"].font = f(12, True, "4FB3F0")
ws[f"{C['Totale']}{R_TOT}"].alignment = DESTRA

r_nota = R_TOT + 2
note = [
    "Note",
    "• Alternative fuori dal totale: pieghevole A3 al posto della brochure, pannelli Pixartprinting al posto di quelli Emmetre. Per scambiarle invertire «Nel totale» sulle due righe.",
    "• Opzionali fuori dal totale base: mini vele, cubi per buca 1 e 18, drone (stimato al massimo di € 400–600). Fotografo: € 150 a giornata per 2 giornate, il massimo di € 80–150.",
    "• Le voci «Da definire» non sono ancora conteggiate. IVA «Esclusa» solo dove il preventivo lo dice: altrimenti da verificare.",
    "• Foto: in Excel 365 selezionare la cella gialla → Inserisci → Immagini → Inserisci nella cella. Le righe libere in fondo sono pronte per nuove voci.",
]
for i, t in enumerate(note):
    cel = ws[f"B{r_nota + i}"]
    cel.value = t
    cel.font = f(9, i == 0, GRIGIO if i else INK)

# ---------------------------------------------------------------- menu e colori di stato
dv_stato = DataValidation(type="list", formula1='"' + ",".join(DATI["stati"].keys()) + '"', allow_blank=True)
dv_si = DataValidation(type="list", formula1='"Sì,No"', allow_blank=True)
dv_ord = DataValidation(type="list", formula1='"Da ordinare,Ordinato,Consegnato,Pagato"', allow_blank=True)
for dv, col in ((dv_stato, "Stato"), (dv_si, "Nel totale"), (dv_ord, "Ordine")):
    ws.add_data_validation(dv)
    dv.add(f"{C[col]}{R0}:{C[col]}{R_LAST}")

COLORI_STATO = {
    "Confermato":  ("E4F2E9", "1F6B40"),
    "Preventivo":  ("E3F0FB", "0A5FA0"),
    "Stima":       ("FFF4D6", "8A6512"),
    "Opzionale":   ("EDEAF7", "4F3F9A"),
    "Alternativa": ("EEF1EE", "5A605A"),
    "Da definire": ("FDE7E4", "A6362F"),
}
area_stato = f"{C['Stato']}{R0}:{C['Stato']}{R_LAST}"
for stato, (sfondo, testo) in COLORI_STATO.items():
    ws.conditional_formatting.add(area_stato, FormulaRule(
        formula=[f'${C["Stato"]}{R0}="{stato}"'], fill=fill(sfondo), font=Font(name=FONT, bold=True, color=testo)))
area_ord = f"{C['Ordine']}{R0}:{C['Ordine']}{R_LAST}"
for stato, (sfondo, testo) in {"Ordinato": ("E3F0FB", "0A5FA0"), "Consegnato": ("E4F2E9", "2C6B4C"),
                               "Pagato": ("D6EDDD", "1F5A3B")}.items():
    ws.conditional_formatting.add(area_ord, FormulaRule(
        formula=[f'${C["Ordine"]}{R0}="{stato}"'], fill=fill(sfondo), font=Font(name=FONT, color=testo)))
# Le voci fuori dal totale si leggono in grigio, così non si sommano a occhio per sbaglio
ws.conditional_formatting.add(f"{C['Totale']}{R0}:{C['Totale']}{R_LAST}", FormulaRule(
    formula=[f'${C["Nel totale"]}{R0}="No"'], font=Font(name=FONT, italic=True, color="9DA39D")))

ws.freeze_panes = f"D{R0}"
ws.auto_filter.ref = f"A{R_HEAD}:{ULTIMA}{R_LAST}"
ws.print_title_rows = f"{R_HEAD}:{R_HEAD}"
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_options.horizontalCentered = True
ws.page_margins.left = ws.page_margins.right = 0.4
ws.oddFooter.left.text = "AIFG · Preventivo Tappa Zero"
ws.oddFooter.right.text = "Pagina &P di &N"

# ---------------------------------------------------------------- riepilogo per area
rs = wb.create_sheet("Riepilogo per area")
rs.sheet_view.showGridLines = False
rs.sheet_properties.tabColor = INK
for col, w in zip("ABCDEFG", (6, 34, 10, 16, 16, 16, 12)):
    rs.column_dimensions[col].width = w
for c in range(1, 8):
    rs.cell(1, c).fill = fill(INK)
    rs.cell(2, c).fill = fill(INK)
rs.row_dimensions[1].height = 28
rs["B1"] = "RIEPILOGO PER AREA"
rs["B1"].font = f(16, True, "FFFFFF")
rs["B2"] = f"{ev['titolo']} · {ev['date']}"
rs["B2"].font = f(10, False, "4FB3F0")

TESTE = ["", "Area", "N. voci", "Totale base", "Opzionali", "Totale", "% base"]
for i, t in enumerate(TESTE, start=1):
    cel = rs.cell(4, i, t.upper())
    cel.font = f(8.5, True, "FFFFFF")
    cel.fill = fill(BLUE)
    cel.alignment = CENTRO if i != 2 else Alignment(vertical="center")
rs.row_dimensions[4].height = 24

P = "Preventivo!"
R_A0, R_AN = 5, 4 + len(DATI["aree"])
R_RT = R_AN + 1  # riga del totale, subito sotto l'ultima area
def prng(col):
    return f"{P}{rng(col)}"

for k, a in enumerate(DATI["aree"]):
    r = 5 + k
    rs.row_dimensions[r].height = 24
    rs.cell(r, 1, a["id"]).font = f(12, True, BLUE)
    rs.cell(r, 1).alignment = CENTRO
    rs.cell(r, 2, a["nome"]).font = f(10.5, True)
    crit = f'{prng("Area")},"{AREE[a["id"]]}"'
    rs.cell(r, 3, f"=COUNTIFS({crit})").alignment = CENTRO
    rs.cell(r, 4, f'=SUMIFS({prng("Totale")},{crit},{prng("Nel totale")},"Sì")').number_format = EUR
    rs.cell(r, 5, f'=SUMIFS({prng("Totale")},{crit},{prng("Stato")},"Opzionale",{prng("Nel totale")},"No")').number_format = EUR
    rs.cell(r, 6, f"=D{r}+E{r}").number_format = EUR
    rs.cell(r, 7, f"=IF($D${R_RT}=0,0,D{r}/$D${R_RT})").number_format = "0.0%"
    for c in range(1, 8):
        rs.cell(r, c).border = Border(bottom=sottile)
        if c > 2:
            rs.cell(r, c).font = f(10.5, c == 4)
            rs.cell(r, c).alignment = DESTRA if c > 3 else CENTRO

rs.row_dimensions[R_RT].height = 26
for c in range(1, 8):
    rs.cell(R_RT, c).fill = fill(INK)
rs[f"B{R_RT}"] = "TOTALE"
rs[f"B{R_RT}"].font = f(10.5, True, "FFFFFF")
for col in "CDEF":
    rs[f"{col}{R_RT}"] = f"=SUM({col}{R_A0}:{col}{R_AN})"
    rs[f"{col}{R_RT}"].font = f(11, True, "4FB3F0" if col == "D" else "FFFFFF")
    rs[f"{col}{R_RT}"].number_format = EUR if col != "C" else "0"
    rs[f"{col}{R_RT}"].alignment = DESTRA if col != "C" else CENTRO
rs[f"G{R_RT}"] = f"=IF($D${R_RT}=0,0,SUM(G{R_A0}:G{R_AN}))"
rs[f"G{R_RT}"].number_format = "0%"
rs[f"G{R_RT}"].font = f(11, True, "FFFFFF")
rs[f"G{R_RT}"].alignment = DESTRA

grafico = BarChart()
grafico.type = "bar"
grafico.style = 10
grafico.title = "Totale base per area (€)"
grafico.y_axis.title = None
grafico.x_axis.title = None
grafico.legend = None
grafico.add_data(Reference(rs, min_col=4, min_row=4, max_row=R_AN), titles_from_data=True)
grafico.set_categories(Reference(rs, min_col=2, min_row=R_A0, max_row=R_AN))
grafico.series[0].graphicalProperties.solidFill = BLUE
grafico.series[0].graphicalProperties.line.solidFill = BLUE
grafico.y_axis.numFmt = '€ #,##0'
grafico.y_axis.majorGridlines = None
grafico.x_axis.scaling.orientation = "maxMin"
grafico.height, grafico.width = 7.5, 18
rs.add_chart(grafico, f"B{R_RT + 3}")

# ================================================================ ATTIVITÀ
wa = wb.create_sheet("Attività", 0)
wb.active = 0
wa.sheet_view.showGridLines = False
wa.sheet_properties.tabColor = "CD212A"

COL_A = [("N.", 5), ("Filone", 22), ("Attività", 34), ("Dettaglio", 48), ("Priorità", 9), ("Tempi", 16),
         ("Responsabili", 24), ("Scadenza", 12), ("Stato", 14), ("Note", 40)]
CA = {nome: get_column_letter(i + 1) for i, (nome, _) in enumerate(COL_A)}
for i, (_, w) in enumerate(COL_A, start=1):
    wa.column_dimensions[get_column_letter(i)].width = w

for r in range(1, 4):
    for c in range(1, len(COL_A) + 1):
        wa.cell(r, c).fill = fill(INK)
wa.row_dimensions[1].height = 30
wa.row_dimensions[2].height = 22
wa.row_dimensions[3].height = 18
logo_a = XLImage(str(DIR / "assets" / "aifg-logo.png"))
logo_a.width, logo_a.height = 64, 63
wa.add_image(logo_a, "A1")
wa["C1"] = f"PIANO OPERATIVO · {ev['titolo'].upper()}"
wa["C1"].font = f(18, True, "FFFFFF")
wa["C1"].alignment = Alignment(vertical="bottom")
wa["C2"] = f"{ev['date']}  ·  {ev['luogo']}, {ev['citta']}"
wa["C2"].font = f(11, False, "4FB3F0")
wa["C3"] = f"Attività istituzionali e preparazione della Tappa Zero  ·  aggiornato al {ev['aggiornato']}"
wa["C3"].font = f(8.5, False, "9DA39D")

ATT = DATI["attivita"]
A_HEAD = 10
A0 = A_HEAD + 1
A_LIBERE = 5
A_LAST = A0 + len(ATT) + A_LIBERE - 1

def arng(col):
    return f"${CA[col]}${A0}:${CA[col]}${A_LAST}"

KPI_A = [
    ("B", "B", "Attività", f'=COUNTA({arng("Attività")})', BLUE),
    ("C", "C", "Priorità 1", f'=COUNTIFS({arng("Priorità")},1)', "CD212A"),
    ("D", "D", "In corso o in revisione", f'=COUNTIFS({arng("Stato")},"In corso")+COUNTIFS({arng("Stato")},"In revisione")', BLUE_DARK),
    ("E", "F", "Da fare", f'=COUNTIFS({arng("Stato")},"Da fare")', GRIGIO),
    ("G", "G", "Fatte", f'=COUNTIFS({arng("Stato")},"Fatto")', "1F6B40"),
]
wa.row_dimensions[5].height = 18
wa.row_dimensions[6].height = 34
for c1, c2, etichetta, formula, colore in KPI_A:
    if c1 != c2:
        wa.merge_cells(f"{c1}5:{c2}5")
        wa.merge_cells(f"{c1}6:{c2}6")
    wa[f"{c1}5"] = etichetta.upper()
    wa[f"{c1}5"].font = f(8, True, GRIGIO)
    wa[f"{c1}5"].alignment = Alignment(vertical="bottom")
    wa[f"{c1}6"] = formula
    wa[f"{c1}6"].number_format = "0"
    wa[f"{c1}6"].font = f(20, True, colore)
    wa[f"{c1}6"].alignment = Alignment(horizontal="left", vertical="center")
    for col in range(wa[f"{c1}6"].column, wa[f"{c2}6"].column + 1):
        wa.cell(6, col).border = Border(bottom=Side(style="medium", color=colore))
        wa.cell(5, col).fill = fill(PANEL)
        wa.cell(6, col).fill = fill(PANEL)

wa["B8"] = "LEGENDA"
wa["B8"].font = f(8, True, GRIGIO)
wa["C8"] = "Priorità 1 = urgente, 2 = media, 3 = bassa, vuota = non assegnata.  Stato: menu a tendina.  Le righe libere in fondo sono pronte per nuove attività."
wa["C8"].font = f(9, False, GRIGIO, italic=True)

wa.row_dimensions[A_HEAD].height = 26
for i, (nome, _) in enumerate(COL_A, start=1):
    cel = wa.cell(A_HEAD, i, nome.upper())
    cel.font = f(8.5, True, "FFFFFF")
    cel.fill = fill(BLUE)
    cel.alignment = CENTRO if nome in ("N.", "Priorità", "Scadenza", "Stato") else Alignment(vertical="center")
    cel.border = Border(bottom=Side(style="medium", color=INK))

GRUPPI = {g["id"]: g["nome"] for g in DATI["gruppiAttivita"]}

def scrivi_a(r, col, val, fmt=None, font=None, align=WRAP, input_=False):
    cel = wa[f"{CA[col]}{r}"]
    cel.value = val
    cel.font = font or f(10)
    cel.alignment = align
    cel.border = BORDO
    if fmt:
        cel.number_format = fmt
    if input_:
        cel.fill = fill(INPUT)
    return cel

for k in range(len(ATT) + A_LIBERE):
    r = A0 + k
    a = ATT[k] if k < len(ATT) else None
    wa.row_dimensions[r].height = 44
    zebra = fill("FFFFFF") if k % 2 == 0 else fill("F8FAF8")
    for col in range(1, len(COL_A) + 1):
        wa.cell(r, col).fill = zebra
    scrivi_a(r, "N.", a["n"] if a else k + 1, font=f(11, True, BLUE), align=CENTRO)
    scrivi_a(r, "Filone", GRUPPI[a["gruppo"]] if a else None, font=f(9, False, GRIGIO))
    scrivi_a(r, "Attività", a["titolo"] if a else None, font=f(10.5, True))
    scrivi_a(r, "Dettaglio", a["dettaglio"] if a else None, font=f(9.5, False, "3A403A"))
    scrivi_a(r, "Priorità", a["priorita"] if a else None, "0", font=f(11, True), align=CENTRO, input_=True)
    scrivi_a(r, "Tempi", (a["tempi"] or None) if a else None, input_=True, font=f(9.5))
    scrivi_a(r, "Responsabili", (", ".join(a["chi"]) or None) if a else None, input_=True, font=f(9.5))
    scad = None
    if a and a["scadenza"]:
        g, m, y = (int(x) for x in a["scadenza"].split("/"))
        from datetime import date
        scad = date(y, m, g)
    scrivi_a(r, "Scadenza", scad, "dd/mm/yyyy", align=CENTRO, input_=True, font=f(10, True, "A6362F"))
    scrivi_a(r, "Stato", a["stato"] if a else None, align=CENTRO, font=f(9.5, True))
    scrivi_a(r, "Note", (a["note"] or None) if a else None, font=f(9, False, "3A403A"))

dv_sa = DataValidation(type="list", formula1='"' + ",".join(DATI["statiAttivita"].keys()) + '"', allow_blank=True)
dv_pr = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
for dv, col in ((dv_sa, "Stato"), (dv_pr, "Priorità")):
    wa.add_data_validation(dv)
    dv.add(f"{CA[col]}{A0}:{CA[col]}{A_LAST}")

COLORI_ATT = {
    "Da fare":      ("EEF1EE", "5A605A"),
    "In corso":     ("E3F0FB", "0A5FA0"),
    "In revisione": ("FFF4D6", "8A6512"),
    "In attesa":    ("EDEAF7", "4F3F9A"),
    "Programmata":  ("EAF3FC", "0A5FA0"),
    "Fatto":        ("E4F2E9", "1F6B40"),
}
for stato, (sfondo, testo) in COLORI_ATT.items():
    wa.conditional_formatting.add(f"{CA['Stato']}{A0}:{CA['Stato']}{A_LAST}", FormulaRule(
        formula=[f'${CA["Stato"]}{A0}="{stato}"'], fill=fill(sfondo), font=Font(name=FONT, bold=True, color=testo)))
for prio, (sfondo, testo) in {1: ("CD212A", "FFFFFF"), 2: ("E8C25A", INK), 3: ("DADDDA", "3A403A")}.items():
    wa.conditional_formatting.add(f"{CA['Priorità']}{A0}:{CA['Priorità']}{A_LAST}", FormulaRule(
        formula=[f'${CA["Priorità"]}{A0}={prio}'], fill=fill(sfondo), font=Font(name=FONT, bold=True, color=testo)))
# Le attività chiuse si leggono attenuate
wa.conditional_formatting.add(f"{CA['Attività']}{A0}:{CA['Attività']}{A_LAST}", FormulaRule(
    formula=[f'${CA["Stato"]}{A0}="Fatto"'], font=Font(name=FONT, bold=True, color="9DA39D", strike=True)))

wa.freeze_panes = f"D{A0}"
wa.auto_filter.ref = f"A{A_HEAD}:{get_column_letter(len(COL_A))}{A_LAST}"
wa.print_title_rows = f"{A_HEAD}:{A_HEAD}"
wa.page_setup.orientation = "landscape"
wa.page_setup.paperSize = wa.PAPERSIZE_A4
wa.page_setup.fitToWidth = 1
wa.page_setup.fitToHeight = 0
wa.sheet_properties.pageSetUpPr.fitToPage = True
wa.oddFooter.left.text = "AIFG · Piano operativo Tappa Zero"
wa.oddFooter.right.text = "Pagina &P di &N"

# ================================================================ COMMUNITY
wc = wb.create_sheet("Community")
wc.sheet_view.showGridLines = False
wc.sheet_properties.tabColor = "4FB3F0"
for col, w in zip("ABCD", (4, 34, 28, 30)):
    wc.column_dimensions[col].width = w
for c in range(1, 5):
    wc.cell(1, c).fill = fill(INK)
    wc.cell(2, c).fill = fill(INK)
wc.row_dimensions[1].height = 28
wc["B1"] = DATI["community"]["nome"].upper()
wc["B1"].font = f(15, True, "FFFFFF")
wc["B2"] = "Struttura dei canali della community ufficiale"
wc["B2"].font = f(10, False, "4FB3F0")
for i, t in enumerate(["", "Canale", "Sezione", "Referente / note"], start=1):
    cel = wc.cell(4, i, t.upper())
    cel.font = f(8.5, True, "FFFFFF")
    cel.fill = fill(BLUE)
r = 5
for c in DATI["community"]["canali"]:
    wc.cell(r, 2, f"{c['ico']}  {c['nome']}").font = f(11, True)
    wc.cell(r, 4).fill = fill(INPUT)
    for col in range(1, 5):
        wc.cell(r, col).border = Border(top=sottile)
    r += 1
    for k, sub in enumerate(c.get("sotto", [])):
        ramo = "└─" if k == len(c["sotto"]) - 1 else "├─"
        wc.cell(r, 3, f"{ramo} {sub}").font = f(10, False, "3A403A")
        wc.cell(r, 4).fill = fill(INPUT)
        r += 1

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"→ {OUT.relative_to(DIR.parent)}")
